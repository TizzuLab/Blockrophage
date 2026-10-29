#!/usr/bin/env python3
"""Draw the Gantt workbook into the course script.

Reads every sheet whose name starts with "Week", resolves merged cells and fill
colours (mapped back to activity blocks via the Legend sheet) and writes:

- stuff/figures/schedule/week-N.svg: one colour-coded calendar per week (days as
  columns, hours as rows, staff and resources rows below the grid),
- the block between <!-- schedule:begin --> and <!-- schedule:end --> in README.md:
  the drawings,
- stuff/tools/schedule/schedule_text.md: a text version per day, which is what the
  companion reads.

Sheet layout (see gantt_build.py): row 1 title, row 2 subtitle, row 3 day headers
(column A: "Time"), then one row per hour with the hour in column A, then a
"Staff" row and a "Resources" row. Cells spanning several hours are merged.

The workbook is the single source for the schedule; the drawings and the README
block are regenerated, never edited by hand.

Usage: python3 stuff/tools/schedule/gantt2readme.py
"""
from __future__ import annotations

import hashlib
import html as _html
import sys
import textwrap
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
XLSX = HERE / "Blockcourse_Gantt_FS2027.xlsx"
README = ROOT / "README.md"
FIG_DIR = ROOT / "stuff" / "figures" / "schedule"
TEXT = HERE / "schedule_text.md"
BEGIN = "<!-- schedule:begin -->"
END = "<!-- schedule:end -->"

HEADER_ROWS = 3          # title, subtitle, day header
FIRST_BODY_ROW = HEADER_ROWS + 1

PAGE_W = 1200            # every week is drawn at the same width so GitHub scales them equally
LABEL_W = 84
HEADER_H = 26
MIN_H = {"hour": 46, "meta": 40}
FONT_SIZE = 12
CHAR_W = FONT_SIZE * 0.56
LINE_H = FONT_SIZE + 3
PAD = 6
THICK = {"1F4E5F", "8B1E3F"}   # plenary and facility: the fixed layer gets a thick border


def rgb6(cell):
    fg = cell.fill.fgColor if cell.fill and cell.fill.fill_type == "solid" else None
    if fg is None or fg.rgb is None or not isinstance(fg.rgb, str):
        return None
    return fg.rgb[-6:].upper()


def font_rgb6(cell):
    try:
        c = cell.font.color.rgb
        return c[-6:] if isinstance(c, str) else "1A1A1A"
    except Exception:
        return "1A1A1A"


def legend_map(wb):
    """fill hex -> short activity label, from the Legend sheet."""
    ws = wb["Legend"]
    m = {}
    for row in ws.iter_rows(min_row=1, max_row=60, max_col=2):
        a, b = row[0], row[1]
        h = rgb6(a)
        if h and a.value and b.value:
            m[h] = str(b.value).split(":")[0].split(" (")[0]
    return m


def legend_items(wb):
    ws = wb["Legend"]
    items = []
    for row in ws.iter_rows(min_row=1, max_row=60, max_col=2):
        a, b = row[0], row[1]
        h = rgb6(a)
        if h and a.value and b.value:
            items.append((str(b.value), h, font_rgb6(a)))
    return items


def body_rows(ws):
    """(row, label, kind) for every labelled row below the header: hours, then staff, resources."""
    out = []
    r = FIRST_BODY_ROW
    while ws.cell(row=r, column=1).value is not None:
        label = " ".join(str(ws.cell(row=r, column=1).value).split())
        kind = "meta" if label.lower().startswith(("staff", "resources", "notes")) else "hour"
        out.append((r, label, kind))
        r += 1
    return out


def blocks_of(ws):
    """(r1, r2, c, cell) for every top-left cell of the body, merged ranges resolved."""
    merged, inside = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng.max_row
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    inside.add((r, c))
    rows = body_rows(ws)
    out = []
    for r, _label, _kind in rows:
        for c in range(2, ws.max_column + 1):
            if (r, c) in inside:
                continue
            out.append((r, merged.get((r, c), r), c, ws.cell(row=r, column=c)))
    return out


def wrap_text(text, width, size=FONT_SIZE, bold=False):
    max_chars = max(6, int((width - 2 * PAD) / (size * (0.62 if bold else 0.56))))
    out = []
    for para in str(text).replace("\r", "").split("\n"):
        para = " ".join(para.split())
        out += textwrap.wrap(para, max_chars, break_long_words=False, break_on_hyphens=False) if para else []
    return out or [""]


def svg_text(x, y, lines, color, bold=False, size=FONT_SIZE):
    out = []
    for i, ln in enumerate(lines):
        out.append(f'<text x="{x:.0f}" y="{y + (i + 1) * (size + 3) - 3:.0f}" font-size="{size}" '
                   f'fill="#{color}"{" font-weight=\"bold\"" if bold else ""}>{_html.escape(ln)}</text>')
    return "\n".join(out)


def week_to_svg(ws, legend_items):
    """One week as an SVG calendar. Row heights grow to fit the text."""
    ndays = ws.max_column - 1
    day_w = (PAGE_W - LABEL_W) / ndays
    rows = body_rows(ws)
    kinds = {r: k for r, _l, k in rows}
    blocks = blocks_of(ws)

    heights = {r: MIN_H[k] for r, _l, k in rows}
    wrapped = {}
    for r1, r2, c, cell in blocks:
        if cell.value is None:
            continue
        lines = wrap_text(cell.value, day_w, bold=bool(cell.font and cell.font.bold))
        wrapped[(r1, c)] = lines
        need = len(lines) * LINE_H + 2 * PAD
        span = r2 - r1 + 1
        per_row = -(-need // span)
        for r in range(r1, r2 + 1):
            heights[r] = max(heights[r], per_row)

    y_of, y = {}, HEADER_H
    for r, _l, _k in rows:
        y_of[r] = y
        y += heights[r]
    legend_h = 26 + 18 * ((len(legend_items) + 1) // 2)
    height = y + 10 + legend_h
    width = PAGE_W

    def x_of(c):
        return LABEL_W + (c - 2) * day_w

    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" width="{width}" height="{height}" '
             f'font-family="Arial, Helvetica, sans-serif">',
             f'<rect x="0" y="0" width="{width}" height="{height}" fill="#ffffff"/>']
    # header row: "Time" and the day names
    parts.append(f'<rect x="0" y="0" width="{LABEL_W}" height="{HEADER_H}" fill="#F2F2F2" stroke="#BFBFBF" stroke-width="0.8"/>')
    parts.append(svg_text(PAD, (HEADER_H - LINE_H) / 2, [str(ws.cell(row=3, column=1).value or "")], "1A1A1A", bold=True))
    for c in range(2, ws.max_column + 1):
        x = x_of(c)
        parts.append(f'<rect x="{x:.1f}" y="0" width="{day_w:.1f}" height="{HEADER_H}" fill="#F2F2F2" stroke="#BFBFBF" stroke-width="0.8"/>')
        txt = " ".join(str(ws.cell(row=3, column=c).value or "").split())
        parts.append(svg_text(x + PAD, (HEADER_H - LINE_H) / 2, [txt], "1A1A1A", bold=True))
    # row labels
    for r, label, _k in rows:
        parts.append(f'<rect x="0" y="{y_of[r]}" width="{LABEL_W}" height="{heights[r]}" fill="#F2F2F2" stroke="#BFBFBF" stroke-width="0.8"/>')
        parts.append(svg_text(PAD, y_of[r] + PAD, [label], "1A1A1A", bold=True))
    # body
    for r1, r2, c, cell in blocks:
        x = x_of(c)
        yy = y_of[r1]
        h = sum(heights[r] for r in range(r1, r2 + 1))
        fill = rgb6(cell) or "FFFFFF"
        thick = fill in THICK
        parts.append(f'<rect x="{x:.1f}" y="{yy}" width="{day_w:.1f}" height="{h}" fill="#{fill}" '
                     f'stroke="#{"1A1A1A" if thick else "BFBFBF"}" stroke-width="{2 if thick else 0.8}"/>')
        if cell.value is not None:
            lines = wrapped[(r1, c)]
            bold = bool(cell.font and cell.font.bold)
            ty = yy + PAD if kinds[r1] == "meta" or r2 == r1 else yy + (h - len(lines) * LINE_H) / 2
            parts.append(svg_text(x + PAD, ty, lines, font_rgb6(cell), bold=bold))
    # legend
    ly = y + 10
    parts.append(f'<text x="4" y="{ly + 14}" font-size="{FONT_SIZE}" font-weight="bold" fill="#1A1A1A">Legend</text>')
    for i, (label, fh, _fc) in enumerate(legend_items):
        col, row = i % 2, i // 2
        lx, lyy = 4 + col * (width // 2), ly + 22 + row * 18
        parts.append(f'<rect x="{lx}" y="{lyy}" width="26" height="13" fill="#{fh}" stroke="#BFBFBF" stroke-width="0.8"/>')
        parts.append(f'<text x="{lx + 32}" y="{lyy + 11}" font-size="11" fill="#1A1A1A">{_html.escape(label)}</text>')
    parts.append("</svg>")
    return "\n".join(parts)


def week_to_md(ws, legend):
    """Text version of one week, one list per day, read by the companion."""
    rows = body_rows(ws)
    label_of = {r: l for r, l, _k in rows}
    kind_of = {r: k for r, _l, k in rows}
    blocks = blocks_of(ws)
    lines = [f"### {ws['A1'].value}", ""]
    for c in range(2, ws.max_column + 1):
        day = " ".join(str(ws.cell(row=3, column=c).value or "").split())
        lines += [f"#### {day}", ""]
        for r1, r2, cc, cell in blocks:
            if cc != c or cell.value is None:
                continue
            text = " ".join(str(cell.value).split())
            tag = legend.get(rgb6(cell))
            tagged = f"[{tag}] {text}" if tag and tag != "Buffer / free" else text
            if kind_of[r1] == "hour":
                start = label_of[r1]
                end_h = int(label_of[r2].split(":")[0]) + 1
                lines.append(f"- {start}-{end_h:02d}:00: {tagged}")
            else:
                lines.append(f"- {label_of[r1]}: {text}")
        lines.append("")
    return lines


def main(xlsx=XLSX, readme=README):
    wb = load_workbook(xlsx)
    legend = legend_map(wb)
    items = legend_items(wb)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    out = [BEGIN,
           f"<!-- Generated from stuff/tools/schedule/{Path(xlsx).name} by stuff/tools/schedule/gantt2readme.py. "
           "Do not edit by hand: edit the workbook and rerun. -->", ""]
    text = ["# Schedule, text version", "",
            f"Generated from `stuff/tools/schedule/{Path(xlsx).name}` by `stuff/tools/schedule/gantt2readme.py` for the",
            "companion; the drawings in README Section 4 are the same content. Activity blocks in brackets;",
            "dark (plenary) blocks are mandatory, everything else is the default plan for all four groups.", ""]
    n_svg = 0
    for ws in wb.worksheets:
        if ws.title.startswith("Week"):
            n = ws.title.split()[-1]
            svg = week_to_svg(ws, items)
            (FIG_DIR / f"week-{n}.svg").write_text(svg, encoding="utf-8")
            n_svg += 1
            # content hash in the link: GitHub and browsers cache images by URL
            digest = hashlib.md5(svg.encode("utf-8")).hexdigest()[:8]
            title = str(ws["A1"].value)
            out += [f"### {title}", "", f"![{title}](stuff/figures/schedule/week-{n}.svg?v={digest})", ""]
            text += week_to_md(ws, legend)
    out.append(END)
    TEXT.write_text("\n".join(text), encoding="utf-8")
    readme_text = Path(readme).read_text(encoding="utf-8")
    a, b = readme_text.index(BEGIN), readme_text.index(END) + len(END)
    Path(readme).write_text(readme_text[:a] + "\n".join(out) + readme_text[b:], encoding="utf-8")
    print(f"wrote {n_svg} drawings to {FIG_DIR}, the schedule block of {readme}, and {TEXT}")


if __name__ == "__main__":
    args = sys.argv[1:]
    main(*(Path(a) for a in args))
