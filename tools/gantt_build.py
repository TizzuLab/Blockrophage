#!/usr/bin/env python3
"""Generate the block-course Gantt template (schedule/Blockcourse_Gantt_FS2027.xlsx).

This is the INITIAL template generator. Once the file has been edited by hand
(groups fill their rows on the planning day), do not rerun this script over it;
run tools/xlsx2md.py instead to refresh schedule/schedule.md from the Excel.

Layout per week sheet
---------------------
- One column block of 8 slots per day: AM 08:30-12:30 and PM 13:30-17:30, one
  cell per hour. A half day is therefore four cells, usually merged.
- Rows: Fixed layer (lectures, facility slots: locked, drawn by staff), one row
  per group with the supervisor's name in the header, a shared Staff row
  (preparation and supervision), a Resources row (hand-written capacity notes)
  and a Notes row.
- Colour = activity block. The checkerboard (two groups wet in the morning, two
  in the afternoon, swapped after lunch) is the default, not a rule.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "schedule" / "Blockcourse_Gantt_FS2027.xlsx"

FONT = "Arial"

# ---------------------------------------------------------------- categories
# key: (legend label, fill hex, font hex, bold)
CATS = {
    "PLENARY": ("Plenary, lecture, debrief (fixed)", "1F4E5F", "FFFFFF", True),
    "FACILITY": ("Facility slot: cytometer, facility lecture (fixed, reserved)", "8B1E3F", "FFFFFF", True),
    "EXP1": ("Wet lab Exp 1: establish the line (growth curve, freeze, re-thaw)", "C6E0B4", "1A1A1A", False),
    "EXP2": ("Wet lab Exp 2: polarize, challenge, read out", "BDD7EE", "1A1A1A", False),
    "DRY": ("Dry lab and desk work: setup, analysis, planning", "FFE699", "1A1A1A", False),
    "POSTER": ("Poster and writing: introduction, significance, figures", "D9C3E9", "1A1A1A", False),
    "EXT": ("Extension / nerd ladder (group's choice)", "F8CBAD", "1A1A1A", False),
    "BUFFER": ("Buffer / free", "EDEDED", "5A5A5A", False),
}

ROWS = [
    ("FIXED", "Fixed layer\n(lectures, facility)"),
    ("A", "Group A\nsupervisor: ________"),
    ("B", "Group B\nsupervisor: ________"),
    ("C", "Group C\nsupervisor: ________"),
    ("D", "Group D\nsupervisor: ________"),
    ("STAFF", "Staff (shared)\nprep and supervision"),
    ("RES", "Resources\n(hand-written capacity)"),
    ("NOTES", "Notes"),
]

SLOTS = ["08:30", "09:30", "10:30", "11:30", "13:30", "14:30", "15:30", "16:30"]
AM = (0, 4)
PM = (4, 4)
ALL = (0, 8)

# ---------------------------------------------------------------- content
# Each week: title, days [(weekday, date)], entries [(row, day_idx, start, length, cat, text)]
# A "pair" helper writes the checkerboard: wet pair gets wet text in the given
# half, dry text in the other half; the other pair the reverse.

def week1():
    days = [("Thu", dt.date(2027, 3, 18)), ("Fri", dt.date(2027, 3, 19))]
    e = []
    # Thursday
    e += [("FIXED", 0, 0, 2, "PLENARY", "Welcome, safety briefing, group formation")]
    e += [("FIXED", 0, 2, 2, "PLENARY", "Lecture: tissue macrophages and nomenclature (recap of Immunology III)")]
    e += [("FIXED", 0, 4, 1, "PLENARY", "Lecture: the techniques of this course, one slide per method")]
    for g in "ABCD":
        e += [(g, 0, 0, 4, "PLENARY", "Plenary")]
        e += [(g, 0, 4, 1, "PLENARY", "Lecture")]
        e += [(g, 0, 5, 3, "DRY", "Experiment planning with your supervisor: fill this Gantt (your row), choose the PRR axis and the extensions")]
    e += [("STAFF", 0, 0, 8, None, "Rooms, safety officer, group list. Supervisors sit with their group from 14:30.")]
    e += [("RES", 0, 0, 8, None, "Seminar room all day")]
    # Friday
    for g in "ABCD":
        e += [(g, 1, 0, 4, "DRY", "Dry lab setup: clone the course repository, build the environment with your LLM, run the ELISA template on example data")]
        e += [(g, 1, 4, 2, "FACILITY", "Facility lecture")]
        e += [(g, 1, 6, 2, "DRY", "Lab tour. Group plans submitted. Read P01 for Monday.")]
    e += [("FIXED", 1, 4, 2, "FACILITY", "Facility lecture: flow cytometry and FlowJo [TODO: confirm slot with the facility]")]
    e += [("STAFF", 1, 0, 8, None, "Dry lab: walk around, unblock installs. Fallback environment ready [TODO]. Collect the group plans at 17:00.")]
    e += [("RES", 1, 0, 8, None, "Computer room or own laptops; Wi-Fi; FlowJo seats [TODO: facility]")]
    return ("Week 1: orientation (Thu 18 - Fri 19 March 2027)", days, e)


def checkerboard(e, day, wet_pair, dry_pair, wet_cat, wet_text, dry_text, am_wet_pair, pm_len=4):
    """Write wet/dry halves for one day. am_wet_pair is wet in the morning.
    pm_len < 4 leaves the last afternoon slot(s) free for a plenary debrief."""
    for g in am_wet_pair:
        e.append((g, day, 0, 4, wet_cat, wet_text))
        e.append((g, day, 4, pm_len, "DRY", dry_text))
    for g in [x for x in "ABCD" if x not in am_wet_pair]:
        e.append((g, day, 0, 4, "DRY", dry_text))
        e.append((g, day, 4, pm_len, wet_cat, wet_text))


def week2():
    days = [("Mon", dt.date(2027, 3, 22)), ("Tue", dt.date(2027, 3, 23)),
            ("Wed", dt.date(2027, 3, 24)), ("Thu", dt.date(2027, 3, 25))]
    e = []
    am = "AB"  # A and B pipette in the morning this week
    wet = [
        "P01 Thaw RAW 264.7, count (T0), seed the flask",
        "P01 Inspect, count (T1), medium change",
        "P02 Count (T2), scrape, split; seed growth-curve wells",
        "P01 Count (T3). P03 Freeze your own vials. Clean-up: nothing stays in culture over the break",
    ]
    dry = [
        "Dry lab: analysis templates on example data (ELISA, Griess)",
        "Dry lab: FlowJo on example data. Poster: draft the introduction",
        "Dry lab: growth curve and doubling time from your counts. Poster: why does this matter (significance)",
        "Dry lab: scRNA-seq part 1, load the preprocessed object, QC figure",
    ]
    for d in range(4):
        checkerboard(e, d, None, None, "EXP1", wet[d], dry[d], am, pm_len=3 if d == 3 else 4)
    e += [("FIXED", 3, 7, 1, "PLENARY", "Debrief week 2: what worked, plan for Exp 2")]
    for g in "ABCD":
        e += [(g, 3, 7, 1, "PLENARY", "Debrief")]
    staff = [
        "Medium warmed, one vial per group at 08:30 and 13:30. Supervisor at the hood with the group for the thaw.",
        "Cultures inspected 08:00. Counting chambers or counter ready.",
        "Plates, scrapers. Look at every split before it goes back in the incubator.",
        "Freezing boxes in the -80. Collect all cultures at 17:00. Easter: thaw backup RAW on Wed 31 Mar, split Fri 2 Apr, so Exp 2 plates can be seeded Mon 5 Apr.",
    ]
    for d, s in enumerate(staff):
        e += [("STAFF", d, 0, 8, None, s)]
    e += [("RES", 0, 0, 8, None, "Hoods: 2 groups per half day [TODO: hoods available]. Incubator: one shelf per group. Cell counter.")]
    e += [("RES", 1, 0, 8, None, "Computer room / laptops; FlowJo seats [TODO]")]
    e += [("RES", 2, 0, 8, None, "Hoods as Mon. Plates: 6-well or 12-well for the growth curve [TODO: lab standard]")]
    e += [("RES", 3, 0, 8, None, "-80 freezer box per group; Mr Frosty or equivalent [TODO: lab standard]")]
    e += [("NOTES", 0, 0, 32, None, "Exp 1 is a real experiment: four counts give a growth curve and a doubling time against the published value; viability of your own vial on Mon 5 Apr is the second readout. Easter break Fri 26 Mar to Sun 4 Apr: no course, no cultures.")]
    return ("Week 2: Exp 1, establish the line (Mon 22 - Thu 25 March 2027)", days, e)


def week3():
    days = [("Mon", dt.date(2027, 4, 5)), ("Tue", dt.date(2027, 4, 6)), ("Wed", dt.date(2027, 4, 7)),
            ("Thu", dt.date(2027, 4, 8)), ("Fri", dt.date(2027, 4, 9))]
    e = []
    am = "CD"  # swapped: C and D pipette in the morning this week
    wet = [
        "P01 Thaw your own vial (viability = Exp 1 readout). P02 Seed Exp 2 plates from the staff culture",
        "P04 Polarize: M(-), M(IFNg), M(IL-4), M(IFNg+LPS), 24 h",
        "P04 Wash. P05 Challenge: LPS + your axis, 24 h. P08 Stain the unchallenged wells, acquire in the facility slot",
        "P05 Harvest supernatants. P06 ELISA IL-6 and TNF. P07 Griess",
        "Plate reader runs, repeats where needed, clean-up. Check your own-vial culture",
    ]
    dry = [
        "Dry lab: scRNA-seq part 2, derive M(IFNg+LPS) and M(IL-4) signatures from the bulk RAW 264.7 dataset",
        "Dry lab: scRNA-seq part 3, score the tissue macrophages. Poster: methods",
        "Dry lab: FlowJo gating of your own data, export statistics",
        "Dry lab: flow statistics in R. Poster: figure 1 draft",
        "Dry lab: ELISA and Griess analysis. Poster: results so far",
    ]
    for d in range(5):
        checkerboard(e, d, None, None, "EXP2", wet[d], dry[d], am, pm_len=3 if d == 4 else 4)
    # Wednesday cytometer slots, one per pair
    e += [("FIXED", 2, 2, 2, "FACILITY", "Cytometer slot, groups C+D [TODO: facility reservation]")]
    e += [("FIXED", 2, 6, 2, "FACILITY", "Cytometer slot, groups A+B [TODO: facility reservation]")]
    e += [("FIXED", 4, 7, 1, "PLENARY", "Debrief Exp 2: results, extension choices, rescue decisions")]
    for g in "ABCD":
        e += [(g, 4, 7, 1, "PLENARY", "Debrief")]
    staff = [
        "Backup RAW thawed Wed 31 Mar, split Fri 2 Apr: enough cells to seed four groups. Plates, cytokines ready.",
        "IFNg, IL-4, LPS aliquots at working concentration. Plate maps checked before anyone pipettes.",
        "Ligands for every axis. Antibody panel and viability dye. Escort the groups to the facility.",
        "ELISA kits, Griess reagents, plate reader booked. Dilution series decided with each group.",
        "Plate reader. -20 space for supernatants. Seed rescue plates Fri PM for groups that need the rescue path (see Legend).",
    ]
    for d, s in enumerate(staff):
        e += [("STAFF", d, 0, 8, None, s)]
    e += [("RES", 0, 0, 8, None, "Hoods: 2 groups per half day. Incubator shelves.")]
    e += [("RES", 1, 0, 8, None, "Hoods. Cytokine stocks [TODO: supplier].")]
    e += [("RES", 2, 0, 8, None, "Cytometer: 2 slots of 2 h [TODO: facility]. Hoods.")]
    e += [("RES", 3, 0, 8, None, "Plate reader (450 nm, 540 nm). Hoods.")]
    e += [("RES", 4, 0, 8, None, "Plate reader. Freezer.")]
    e += [("NOTES", 0, 0, 40, None, "One 24 h challenge arm only; the 6 h RNA arm belongs to the qPCR extension X01. A group that starts a step in the morning harvests it the next morning; pairs stay fixed within the week.")]
    return ("Week 3: Exp 2, polarize and challenge (Mon 5 - Fri 9 April 2027)", days, e)


def week4():
    days = [("Mon", dt.date(2027, 4, 12)), ("Tue", dt.date(2027, 4, 13)), ("Wed", dt.date(2027, 4, 14)),
            ("Thu", dt.date(2027, 4, 15)), ("Fri", dt.date(2027, 4, 16))]
    e = []
    am = "AB"
    wet = [
        "Extension of your choice, or rescue path for Exp 2 (see Legend)",
        "Extension, continued",
    ]
    dry = [
        "Dry lab: analysis of Exp 2. Poster: figures",
        "Dry lab: extension analysis. Poster: figures and legends",
    ]
    for d in range(2):
        checkerboard(e, d, None, None, "EXT", wet[d], dry[d], am)
    e += [("FIXED", 1, 7, 1, "PLENARY", "Student-organised event (evening, from 17:30)")]
    # Wednesday: writing day
    e += [("FIXED", 2, 0, 4, "PLENARY", "Scientific writing day: from data to argument, poster structure, figure design, honest limits")]
    e += [("FIXED", 2, 6, 2, "PLENARY", "Failure round: each group presents one thing that went wrong and what it taught them")]
    for g in "ABCD":
        e += [(g, 2, 0, 4, "PLENARY", "Writing day")]
        e += [(g, 2, 4, 2, "POSTER", "Drafting with your supervisor")]
        e += [(g, 2, 6, 2, "PLENARY", "Failure round")]
    # Thursday: poster production
    e += [("FIXED", 3, 6, 2, "PLENARY", "Feedback round with the teaching staff. Nerd ladder closes 17:00")]
    for g in "ABCD":
        e += [(g, 3, 0, 6, "POSTER", "Poster production")]
        e += [(g, 3, 6, 2, "PLENARY", "Feedback round")]
    # Friday
    e += [("FIXED", 4, 0, 4, "PLENARY", "Poster presentations and discussion")]
    e += [("FIXED", 4, 4, 4, "PLENARY", "Prize, course evaluation, lab clean-up")]
    for g in "ABCD":
        e += [(g, 4, 0, 8, "PLENARY", "Presentations, prize, evaluation, clean-up")]
    staff = [
        "Reagents for the extensions agreed on Fri 9 Apr. Rescue plates: polarize Mon.",
        "Event logistics with the organisers. Rescue plates: challenge Tue.",
        "Poster template distributed, printing booked [TODO]. Rescue plates: harvest 08:30.",
        "Printing. Rescue groups: ELISA in the morning, result on the poster by 17:00.",
        "Jury, prize, evaluation form, clean-up checklist.",
    ]
    for d, s in enumerate(staff):
        e += [("STAFF", d, 0, 8, None, s)]
    e += [("RES", 0, 0, 16, None, "Hoods Mon and Tue only. Extension-specific instruments as agreed.")]
    e += [("RES", 2, 0, 24, None, "Seminar room Wed to Fri. Printer [TODO]. Poster boards.")]
    e += [("NOTES", 0, 0, 40, None, "A full repeat of Exp 2 does not fit in week 4. Rescue path for a failed Exp 2: staff seed plates Fri 9 Apr, group polarizes Mon, challenges Tue, harvests Wed 08:30, ELISA Thu morning.")]
    return ("Week 4: extensions, writing, poster (Mon 12 - Fri 16 April 2027)", days, e)


# ---------------------------------------------------------------- rendering
thin = Side(style="thin", color="BFBFBF")
thick = Side(style="medium", color="1A1A1A")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FIXED_BORDER = Border(left=thick, right=thick, top=thick, bottom=thick)


def style_range(ws, r1, c1, r2, c2, fill=None, font=None, border=BORDER, align=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if border:
                cell.border = border
            if align:
                cell.alignment = align


def render_week(wb, title, days, entries):
    sheet_title = title.split(":")[0]
    ws = wb.create_sheet(sheet_title)
    ncols = 1 + 8 * len(days)
    ws.column_dimensions["A"].width = 24
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1F4E5F")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A2"] = "Fixed layer is drawn by staff and locked. Groups fill their own row with their supervisor on the planning day. Colours: see Legend."
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="5A5A5A")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    hdr_fill = PatternFill("solid", fgColor="F2F2F2")
    hdr_font = Font(name=FONT, size=10, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # row 3: days, row 4: AM/PM, row 5: slots
    for d, (wd, date) in enumerate(days):
        c0 = 2 + 8 * d
        ws.cell(row=3, column=c0, value=f"{wd} {date.day} {date.strftime('%b')}")
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 7)
        ws.cell(row=4, column=c0, value="Morning 08:30-12:30")
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 3)
        ws.cell(row=4, column=c0 + 4, value="Afternoon 13:30-17:30")
        ws.merge_cells(start_row=4, start_column=c0 + 4, end_row=4, end_column=c0 + 7)
        for s, t in enumerate(SLOTS):
            ws.cell(row=5, column=c0 + s, value=t)
        style_range(ws, 3, c0, 5, c0 + 7, fill=hdr_fill, font=hdr_font, align=center)
    ws.cell(row=3, column=1, value="Day").font = hdr_font
    ws.cell(row=4, column=1, value="Half day").font = hdr_font
    ws.cell(row=5, column=1, value="Slot").font = hdr_font
    style_range(ws, 3, 1, 5, 1, fill=hdr_fill, font=hdr_font, align=center)

    row_index = {}
    r = 6
    for key, label in ROWS:
        row_index[key] = r
        ws.cell(row=r, column=1, value=label)
        ws.row_dimensions[r].height = 30 if key == "NOTES" else (48 if key in ("STAFF", "RES") else 78)
        style_range(ws, r, 1, r, ncols, border=BORDER,
                    align=Alignment(vertical="top", wrap_text=True))
        ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
        ws.cell(row=r, column=1).fill = hdr_fill
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        if key in ("STAFF", "RES", "NOTES"):
            style_range(ws, r, 2, r, ncols, fill=PatternFill("solid", fgColor="FAFAFA"))
        r += 1

    # buffer fill for group cells first, entries overwrite
    for g in "ABCD":
        style_range(ws, row_index[g], 2, row_index[g], ncols,
                    fill=PatternFill("solid", fgColor=CATS["BUFFER"][1]))

    for (rowkey, day, start, length, cat, text) in entries:
        rr = row_index[rowkey]
        c1 = 2 + 8 * day + start
        c2 = min(c1 + length - 1, ncols)
        cell = ws.cell(row=rr, column=c1, value=text)
        if cat:
            _, fill_hex, font_hex, bold = CATS[cat]
            fill = PatternFill("solid", fgColor=fill_hex)
            font = Font(name=FONT, size=9, bold=bold, color=font_hex)
            border = FIXED_BORDER if cat in ("PLENARY", "FACILITY") else BORDER
        else:
            fill, font, border = None, Font(name=FONT, size=9), BORDER
        style_range(ws, rr, c1, rr, c2, fill=fill, font=font, border=border,
                    align=Alignment(vertical="top", wrap_text=True))
        if c2 > c1:
            ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
        cell.value = text

    ws.freeze_panes = "B6"
    ws.sheet_view.zoomScale = 80
    return ws


def render_legend(wb):
    ws = wb.active
    ws.title = "Legend"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Macrophage block course FS 2027: schedule template"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1F4E5F")
    ws["A2"] = f"Template generated {dt.date.today().isoformat()} by tools/gantt_build.py. After hand editing, refresh schedule.md with tools/xlsx2md.py; do not rerun the generator."
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="5A5A5A")

    r = 4
    ws.cell(row=r, column=1, value="Colour").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value="Activity block").font = Font(name=FONT, bold=True)
    r += 1
    for key, (label, fill_hex, font_hex, bold) in CATS.items():
        a = ws.cell(row=r, column=1, value=key)
        a.fill = PatternFill("solid", fgColor=fill_hex)
        a.font = Font(name=FONT, size=10, bold=bold, color=font_hex)
        a.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=label).font = Font(name=FONT, size=10)
        r += 1

    r += 1
    rules = [
        ("How to read it", "One sheet per week. Eight cells per day: four in the morning, four in the afternoon, one per hour. Rows: fixed layer, one row per group, shared staff row, resources, notes."),
        ("Fixed layer", "Lecture slots, facility lecture, cytometer slots, writing day, poster day. Drawn by the teaching staff, thick border, locked. Everything else is flexible."),
        ("Group rows", "Each group fills its own row with its supervisor on the planning day (Thu 18 March, afternoon) and submits the plan on Friday. The plan is the group's first deliverable. The supervisor travels with the group for the whole course and may fold in experiments the lab needs anyway; those go into the group notebook like any other experiment."),
        ("Checkerboard default", "Two groups at the bench in the morning and two in the dry lab, swapped after lunch. It is a default, not a rule: it keeps two hoods sufficient and the cytometer at two groups per slot. A group that deviates justifies its hood time against the other three. Pairs stay fixed within a week, because the group that starts a 24 h step harvests it. The morning pair swaps between weeks."),
        ("Core experiments", "Exp 1 (week 2): establish the RAW 264.7 line. Thaw, daily counts (growth curve, doubling time), split, seed, freeze your own vials; viability of your vial after the break is the second readout. Exp 2 (week 3): polarize M(-), M(IFNg), M(IL-4), M(IFNg+LPS); challenge with LPS and your PRR axis, 24 h; ELISA, Griess, flow cytometry. Protocol IDs (P01...) refer to protocols/ in the course repository."),
        ("Dry lab", "Week 1: environment and templates. Week 2: templates on example data, growth curve, scRNA-seq part 1. Week 3: signatures from the bulk RAW 264.7 dataset, scoring of the tissue macrophages, own flow and ELISA data. Week 4: extension analysis and poster figures. Downtime at any point: poster introduction and significance."),
        ("Rescue path", "A full repeat of Exp 2 does not fit in week 4. A group whose Exp 2 failed takes the rescue path: staff seed plates Fri 9 April, polarize Mon 12, challenge Tue 13, harvest Wed 14 at 08:30, ELISA Thu 15 morning."),
        ("Easter break", "Fri 26 March to Sun 4 April. No course, no cultures. Staff thaw backup RAW on Wed 31 March and split Fri 2 April [TODO: confirm dates with the department]."),
    ]
    for k, v in rules:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * (1 + len(v) // 110)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Open items").font = Font(name=FONT, size=10, bold=True, color="8B1E3F")
    r += 1
    todos = [
        "Cytometer slots: dates and times from the facility (week 3 Wednesday, two slots; extension slots week 4).",
        "Facility lecture: slot on Fri 19 March.",
        "FlowJo seats: number, and whether they run on student laptops or facility workstations (decides where the dry half of the day happens).",
        "Hoods available to the course per half day.",
        "Supervisor names in the group row headers.",
        "qPCR: extension (current plan) or core.",
        "Course dates and Easter break: confirm with the department.",
        "Poster printing: where, deadline.",
    ]
    for t in todos:
        ws.cell(row=r, column=2, value="[TODO] " + t).font = Font(name=FONT, size=10)
        r += 1
    return ws


def render_overview(wb, weeks):
    ws = wb.create_sheet("Overview", 1)
    ws.column_dimensions["A"].width = 14
    for c in range(2, 7):
        ws.column_dimensions[get_column_letter(c)].width = 34
    ws["A1"] = "Overview: the four weeks at a glance (default plan)"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1F4E5F")
    hdr = ["", "Mon", "Tue", "Wed", "Thu", "Fri"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.alignment = Alignment(horizontal="center")
    summary = {
        "Week 1": {"Thu": ("PLENARY", "Welcome, lectures, planning with supervisor"), "Fri": ("DRY", "Dry lab setup; facility lecture; plans submitted")},
        "Week 2": {"Mon": ("EXP1", "Exp 1: thaw, T0"), "Tue": ("EXP1", "Exp 1: T1, medium change"), "Wed": ("EXP1", "Exp 1: T2, split, seed"), "Thu": ("EXP1", "Exp 1: T3, freeze; debrief")},
        "Easter": {"Mon": ("BUFFER", "Break 26 Mar - 4 Apr; staff thaw backup cells")},
        "Week 3": {"Mon": ("EXP2", "Re-thaw own vial; seed Exp 2"), "Tue": ("EXP2", "Polarize"), "Wed": ("EXP2", "Challenge; flow (facility slots)"), "Thu": ("EXP2", "Harvest; ELISA, Griess"), "Fri": ("EXP2", "Readouts, analysis; debrief")},
        "Week 4": {"Mon": ("EXT", "Extensions / rescue path"), "Tue": ("EXT", "Extensions; evening event"), "Wed": ("PLENARY", "Writing day; failure round"), "Thu": ("POSTER", "Poster production; feedback"), "Fri": ("PLENARY", "Presentations, prize")},
    }
    r = 4
    for wk, cells in summary.items():
        ws.cell(row=r, column=1, value=wk).font = Font(name=FONT, bold=True)
        ws.row_dimensions[r].height = 36
        for c, day in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri"], start=2):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if day in cells:
                cat, text = cells[day]
                cell.value = text
                _, fh, fc, b = CATS[cat]
                cell.fill = PatternFill("solid", fgColor=fh)
                cell.font = Font(name=FONT, size=10, bold=b, color=fc)
        r += 1
    ws.cell(row=r + 1, column=1, value="Every day in weeks 2 to 4 (Mon-Tue) is half wet lab, half dry lab per group; see the week sheets.").font = Font(name=FONT, size=9, italic=True, color="5A5A5A")
    return ws


def main():
    wb = Workbook()
    render_legend(wb)
    weeks = [week1(), week2(), week3(), week4()]
    for title, days, entries in weeks:
        render_week(wb, title, days, entries)
    render_overview(wb, weeks)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
