#!/usr/bin/env python3
"""Draw the hand-edited Gantt workbook into the course script.

Reads every sheet whose name starts with "Week", resolves merged cells and fill
colours (mapped back to activity blocks via the Legend sheet) and writes:

- figures/schedule/week-N.svg: one colour-coded drawing per week,
- the block between <!-- schedule:begin --> and <!-- schedule:end --> in README.md:
  the drawings plus a folded text table per day, which is what the LLM companion
  reads.

The workbook is the single source for the schedule; the drawings and the README
block are regenerated, never edited by hand.

Usage: python3 tools/schedule/gantt2readme.py
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
XLSX = HERE / "Blockcourse_Gantt_FS2027.xlsx"
README = ROOT / "README.md"
FIG_DIR = ROOT / "figures" / "schedule"
BEGIN = "<!-- schedule:begin -->"
END = "<!-- schedule:end -->"

HEADER_ROWS = 5  # title, subtitle, day, half day, slot


def rgb6(cell):
    fg = cell.fill.fgColor if cell.fill and cell.fill.fill_type == "solid" else None
    if fg is None or fg.rgb is None or not isinstance(fg.rgb, str):
        return None
    return fg.rgb[-6:].upper()


def legend_map(wb):
    """fill hex -> short activity label, from the Legend sheet."""
    ws = wb["Legend"]
    m = {}
    for row in ws.iter_rows(min_row=1, max_row=60, max_col=2):
        a, b = row[0], row[1]
        h = rgb6(a)
        if h and a.value and b.value:
            m[h] = str(b.value).split(":")[0].split(" (")[0]
    return m


def merged_lookup(ws):
    """cell coordinate -> top-left anchor coordinate for merged ranges."""
    anchor = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(row=rng.min_row, column=rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                anchor[(r, c)] = tl
    return anchor


def collect(ws, anchor, legend, row, c_from, c_to):
    """Unique (label, text) pairs across a span of columns, in order."""
    out, seen = [], set()
    for c in range(c_from, c_to + 1):
        cell = anchor.get((row, c), ws.cell(row=row, column=c))
        if cell.value is None:
            continue
        key = cell.coordinate
        if key in seen:
            continue
        seen.add(key)
        label = legend.get(rgb6(cell))
        text = " ".join(str(cell.value).split())
        out.append(f"[{label}] {text}" if label and label != "Buffer / free" else text)
    return "; ".join(out) if out else "-"


import html as _html
import textwrap

PAGE_W = 1200          # every week is drawn at the same width so GitHub scales them equally
LABEL_W = 150
HEADER_H = (26, 22, 18)
MIN_ROW_H = {"fixed": 40, "group": 44, "staff": 40, "resources": 40, "notes": 36}
FONT_SIZE = 12
CHAR_W = FONT_SIZE * 0.56
LINE_H = FONT_SIZE + 3
PAD = 6


def row_kind(label):
    l = label.lower()
    if l.startswith("fixed"): return "fixed"
    if l.startswith("group"): return "group"
    if l.startswith("staff"): return "staff"
    if l.startswith("resources"): return "resources"
    if l.startswith("notes"): return "notes"
    return "group"


def font_rgb6(cell):
    try:
        c = cell.font.color.rgb
        return c[-6:] if isinstance(c, str) else "1A1A1A"
    except Exception:
        return "1A1A1A"


def wrap_text(text, width, size=FONT_SIZE):
    max_chars = max(6, int((width - 2 * PAD) / (size * 0.56)))
    out = []
    for para in str(text).replace("\r", "").split("\n"):
        para = " ".join(para.split())
        out += textwrap.wrap(para, max_chars, break_long_words=False, break_on_hyphens=False) if para else []
    return out or [""]


NARROW_W = 70    # blocks narrower than this get rotated text when it does not fit horizontally
ROT_MAX_H = 280  # rotated text never asks for more than this height; it shrinks, then truncates


def wrap_rotated(text, width):
    """Wrap for text running bottom-to-top in a narrow block. The block width limits
    the number of lines; returns (lines, needed_height, font_size)."""
    text = " ".join(str(text).split())
    for size in (FONT_SIZE, 10):
        lh, cw = size + 3, size * 0.56
        max_lines = max(1, int((width - 2 * PAD) / lh))
        L = max(6, -(-len(text) // max_lines))
        while True:
            lines = textwrap.wrap(text, L, break_long_words=False, break_on_hyphens=False) or [""]
            if len(lines) <= max_lines:
                break
            L += 2
        need = max(len(ln) for ln in lines) * cw + 2 * PAD
        if need <= ROT_MAX_H:
            return lines, need, size
    maxc = int((ROT_MAX_H - 2 * PAD) / cw)
    lines = [ln if len(ln) <= maxc else ln[: maxc - 1] + "…" for ln in lines]
    return lines, ROT_MAX_H, size


def svg_text_rotated(x, y_bottom, lines, color, bold=False, size=FONT_SIZE):
    out = []
    for i, ln in enumerate(lines):
        tx = x + PAD + (i + 1) * (size + 3) - 3
        out.append(f'<text x="{tx:.0f}" y="{y_bottom:.0f}" transform="rotate(-90 {tx:.0f} {y_bottom:.0f})" font-size="{size}" '
                   f'fill="#{color}"{" font-weight=\"bold\"" if bold else ""}>{_html.escape(ln)}</text>')
    return "\n".join(out)


def svg_text(x, y, lines, color, bold=False, size=FONT_SIZE):
    out = []
    for i, ln in enumerate(lines):
        out.append(f'<text x="{x:.0f}" y="{y + (i + 1) * (size + 3) - 3:.0f}" font-size="{size}" '
                   f'fill="#{color}"{" font-weight=\"bold\"" if bold else ""}>{_html.escape(ln)}</text>')
    return "\n".join(out)


def week_to_svg(ws, legend_items):
    """Draw one week sheet as an SVG string. legend_items: [(label, fill hex, font hex)].

    Every week is drawn at PAGE_W so the font renders at the same size on GitHub
    regardless of how many days the week has. Row heights grow to fit the text.
    Adjacent group rows that hold the same block (same span, text and colour) are
    merged into one cell so a plenary or a shared plan is written once.
    """
    merged, inside = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    inside.add((r, c))
    ncols = ws.max_column
    nslots = ncols - 1
    slot_w = (PAGE_W - LABEL_W) / nslots
    width = PAGE_W
    show_slots = slot_w >= 40          # hourly time row only when it is legible

    def x_of(c):
        return 0 if c == 1 else LABEL_W + (c - 2) * slot_w

    def cell_w(c1, c2):
        return (x_of(c2) + (slot_w if c2 > 1 else LABEL_W)) - x_of(c1)

    labels = []
    r = HEADER_ROWS + 1
    while ws.cell(row=r, column=1).value:
        labels.append((r, " ".join(str(ws.cell(row=r, column=1).value).split())))
        r += 1
    kinds = {rr: row_kind(l) for rr, l in labels}
    body_rows = [rr for rr, _ in labels]

    # blocks: (r1, r2, c1, c2, cell) with vertical merging of identical group cells
    blocks, taken = [], set()
    for rr in body_rows:
        for c in range(1, ncols + 1):
            if (rr, c) in inside or (rr, c) in taken:
                continue
            c2 = merged.get((rr, c), (rr, c))[1]
            cell = ws.cell(row=rr, column=c)
            r2 = rr
            if c > 1 and kinds[rr] == "group" and cell.value is not None:
                key = (c2, " ".join(str(cell.value).split()), rgb6(cell))
                nxt = rr + 1
                while nxt in kinds and kinds[nxt] == "group" and (nxt, c) not in inside:
                    ncell = ws.cell(row=nxt, column=c)
                    nc2 = merged.get((nxt, c), (nxt, c))[1]
                    if ncell.value is None or (nc2, " ".join(str(ncell.value).split()), rgb6(ncell)) != key:
                        break
                    r2 = nxt
                    for cc in range(c, c2 + 1):
                        taken.add((nxt, cc))
                    nxt += 1
            blocks.append((rr, r2, c, c2, cell))

    # row heights: fit the tallest wrapped text in each row (shared across a vertical span)
    heights = {rr: MIN_ROW_H[kinds[rr]] for rr in body_rows}
    wrapped = {}
    for (r1, r2, c1, c2, cell) in blocks:
        if cell.value is None:
            continue
        w = cell_w(c1, c2)
        lines = wrap_text(cell.value, w)
        overflow = max(len(ln) for ln in lines) * CHAR_W > w - 2 * PAD
        rotated = c1 > 1 and ((w < NARROW_W and len(lines) > 2) or (overflow and w < 110))
        size = FONT_SIZE
        if rotated:
            lines, need, size = wrap_rotated(cell.value, w)
        else:
            need = len(lines) * LINE_H + 2 * PAD
        wrapped[(r1, c1)] = (lines, rotated, size)
        span = r2 - r1 + 1
        per_row = -(-need // span)
        for rr in range(r1, r2 + 1):
            heights[rr] = max(heights[rr], per_row)

    header_rows = [3, 4, 5] if show_slots else [3, 4]
    y_of, y = {}, 0
    for i, rr in enumerate(header_rows):
        y_of[rr] = y; y += HEADER_H[rr - 3]
    for rr in body_rows:
        y_of[rr] = y; y += heights[rr]
    legend_h = 26 + 18 * ((len(legend_items) + 1) // 2)
    height = y + 10 + legend_h

    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" width="{width}" height="{height}" '
             f'font-family="Arial, Helvetica, sans-serif">',
             f'<rect x="0" y="0" width="{width}" height="{height}" fill="#ffffff"/>']
    # headers
    for rr in header_rows:
        rh = HEADER_H[rr - 3]
        for c in range(1, ncols + 1):
            if (rr, c) in inside:
                continue
            c2 = merged.get((rr, c), (rr, c))[1]
            cell = ws.cell(row=rr, column=c)
            x, w = x_of(c), cell_w(c, c2)
            parts.append(f'<rect x="{x:.1f}" y="{y_of[rr]}" width="{w:.1f}" height="{rh}" fill="#F2F2F2" stroke="#BFBFBF" stroke-width="0.8"/>')
            if cell.value is not None:
                size = 10 if rr == 5 else FONT_SIZE
                txt = " ".join(str(cell.value).split())
                if len(txt) * size * 0.56 > w - 2 * PAD:
                    txt = txt.split()[0]           # "Morning 09:00-12:00" -> "Morning"
                if len(txt) * size * 0.56 > w - 2 * PAD and size > 9:
                    size = 9
                if len(txt) * size * 0.56 > w - 2 * PAD:
                    txt = ""                        # too narrow: the body cells carry the label
                if txt:
                    parts.append(svg_text(x + PAD, y_of[rr] + (rh - (size + 3)) / 2, [txt], "1A1A1A", bold=True, size=size))
    # body
    for (r1, r2, c1, c2, cell) in blocks:
        x, w = x_of(c1), cell_w(c1, c2)
        yy = y_of[r1]
        h = sum(heights[rr] for rr in range(r1, r2 + 1))
        fill = rgb6(cell) or ("F2F2F2" if c1 == 1 else "FFFFFF")
        bold = bool(cell.font and cell.font.bold)
        thick = fill in ("1F4E5F", "8B1E3F")
        parts.append(f'<rect x="{x:.1f}" y="{yy}" width="{w:.1f}" height="{h}" fill="#{fill}" '
                     f'stroke="#{"1A1A1A" if thick else "BFBFBF"}" stroke-width="{2 if thick else 0.8}"/>')
        if cell.value is not None:
            lines, rotated, size = wrapped[(r1, c1)]
            if rotated:
                tlen = max(len(ln) for ln in lines) * size * 0.56
                parts.append(svg_text_rotated(x, yy + (h + tlen) / 2, lines, font_rgb6(cell), bold=bold, size=size))
            else:
                ty = yy + PAD if r2 == r1 else yy + (h - len(lines) * LINE_H) / 2
                parts.append(svg_text(x + PAD, ty, lines, font_rgb6(cell), bold=bold or c1 == 1))
    # legend
    ly = y + 10
    parts.append(f'<text x="4" y="{ly + 14}" font-size="{FONT_SIZE}" font-weight="bold" fill="#1A1A1A">Legend</text>')
    for i, (label, fh, fc) in enumerate(legend_items):
        col, row = i % 2, i // 2
        lx, lyy = 4 + col * (width // 2), ly + 22 + row * 18
        parts.append(f'<rect x="{lx}" y="{lyy}" width="26" height="13" fill="#{fh}" stroke="#BFBFBF" stroke-width="0.8"/>')
        parts.append(f'<text x="{lx + 32}" y="{lyy + 11}" font-size="11" fill="#1A1A1A">{_html.escape(label)}</text>')
    parts.append("</svg>")
    return "\n".join(parts)


def legend_items(wb):
    ws = wb["Legend"]
    items = []
    for row in ws.iter_rows(min_row=1, max_row=60, max_col=2):
        a, b = row[0], row[1]
        h = rgb6(a)
        if h and a.value and b.value:
            items.append((str(b.value), h, font_rgb6(a)))
    return items


def week_to_md(ws, legend, svg_name=None):
    anchor = merged_lookup(ws)
    lines = [f"### {ws['A1'].value}", ""]
    if svg_name:
        lines += [f"![{ws['A1'].value}]({svg_name})", "", "<details><summary>Text version (what the companion reads)</summary>", ""]
    # days from row 3 merged headers
    days = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 3 and rng.min_col >= 2:
            days.append((rng.min_col, ws.cell(row=3, column=rng.min_col).value))
    days.sort()
    labels = []
    r = HEADER_ROWS + 1
    while ws.cell(row=r, column=1).value:
        labels.append((r, " ".join(str(ws.cell(row=r, column=1).value).split())))
        r += 1
    notes_rows = [r for r, l in labels if l.lower().startswith("notes")]
    # per-day segments (morning, lunch, afternoon ...) from the merged ranges of row 4
    seg_of = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 4 and rng.min_col >= 2:
            seg_of.setdefault(rng.min_col, (rng.min_col, rng.max_col, ws.cell(row=4, column=rng.min_col).value))
    for c in range(2, ws.max_column + 1):
        if c not in seg_of and ws.cell(row=4, column=c).value and c not in {k for k in seg_of}:
            if not any(a <= c <= b for a, b, _ in seg_of.values()):
                seg_of[c] = (c, c, ws.cell(row=4, column=c).value)
    for i, (c0, day) in enumerate(days):
        c_end = days[i + 1][0] - 1 if i + 1 < len(days) else ws.max_column
        segs = sorted(v for k, v in seg_of.items() if c0 <= k <= c_end)
        hdr = " | ".join(" ".join(str(h).split()) for _, _, h in segs)
        lines += [f"#### {day}", "", f"| | {hdr} |", "|" + "---|" * (len(segs) + 1)]
        for rr, label in labels:
            if rr in notes_rows:
                continue
            cells, prev = [], None
            for a, b, _ in segs:
                txt = collect(ws, anchor, legend, rr, a, b)
                cells.append("(same as previous column)" if txt != "-" and txt == prev else txt)
                prev = txt
            lines.append(f"| {label} | " + " | ".join(cells) + " |")
        lines.append("")
    for rr in notes_rows:
        note = collect(ws, anchor, legend, rr, 2, ws.max_column)
        if note != "-":
            lines += [f"Notes: {note}", ""]
    if svg_name:
        lines += ["</details>", ""]
    return lines


def main(xlsx=XLSX, readme=README):
    wb = load_workbook(xlsx)
    legend = legend_map(wb)
    items = legend_items(wb)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    out = [BEGIN,
           f"<!-- Generated from tools/schedule/{Path(xlsx).name} by tools/schedule/gantt2readme.py. "
           "Do not edit by hand: edit the workbook and rerun. -->", ""]
    n_svg = 0
    for ws in wb.worksheets:
        if ws.title.startswith("Week"):
            n = ws.title.split()[-1]
            svg = week_to_svg(ws, items)
            (FIG_DIR / f"week-{n}.svg").write_text(svg, encoding="utf-8")
            n_svg += 1
            # content hash in the link: GitHub and browsers cache images by URL
            digest = hashlib.md5(svg.encode("utf-8")).hexdigest()[:8]
            out += week_to_md(ws, legend, f"figures/schedule/week-{n}.svg?v={digest}")
    out.append(END)
    text = Path(readme).read_text(encoding="utf-8")
    a, b = text.index(BEGIN), text.index(END) + len(END)
    Path(readme).write_text(text[:a] + "\n".join(out) + text[b:], encoding="utf-8")
    print(f"wrote {n_svg} drawings to {FIG_DIR} and the schedule block of {readme}")


if __name__ == "__main__":
    args = sys.argv[1:]
    main(*(Path(a) for a in args))
